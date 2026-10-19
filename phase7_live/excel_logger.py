# -*- coding: utf-8 -*-
"""
phase7_live/excel_logger.py
===========================
Daily Excel Logger for Phase 7 Live Inspection.
Maintains running counts (pass, fail, total) and records inspection rows.
Features:
- Embeds defect inspection image into the Excel row whenever a part FAILS!
- Highlights individual missing/defective features with soft red fill and bold dark red text.
- Formats overall PASS with soft green, and FAIL with bold red.
- Detailed Failure Reason column.
"""

import os
import csv
from datetime import datetime
from pathlib import Path
import cv2

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.drawing.image import Image as OpenpyxlImage
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

ROOT = Path(__file__).resolve().parent.parent

class DailyExcelLogger:
    """Manages daily inspection logs, defect image capture, and production counts."""

    def __init__(self, log_dir: str | Path | None = None):
        if log_dir is None:
            log_dir = ROOT / "outputs" / "reports"
        self.output_dir = Path(log_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)
        
        # Dedicated folder to store captured defect thumbnail images
        self.defect_img_dir = self.output_dir / "defect_captures"
        self.defect_img_dir.mkdir(parents=True, exist_ok=True)

        self.current_date = datetime.now().strftime("%Y-%m-%d")
        self.filepath = self.output_dir / f"inspection_log_{self.current_date}.xlsx"
        self.csv_path = self.output_dir / f"inspection_log_{self.current_date}.csv"

        self.headers = [
            "Inspection #",
            "Timestamp",
            "Device",
            "Overall Result",
            "Failure Reason",
            "Defect Image",
            "Spline",
            "Top Nib",
            "Bottom Nib",
            "Hole TL",
            "Hole TR",
            "Hole BL",
            "Hole BR",
            "Min Confidence %",
            "Reproj Error (px)"
        ]

        self.total_pass = 0
        self.total_fail = 0
        self.total_inspections = 0

        self._init_files()

    def _init_files(self):
        if HAS_OPENPYXL and not self.filepath.exists():
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Inspection Summary"
            
            # Title
            ws.merge_cells("A1:O1")
            ws["A1"] = f"DAILY PRODUCTION LOG — {self.current_date}"
            ws["A1"].font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
            ws["A1"].fill = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")
            ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[1].height = 30
            
            # Headers
            header_fill = PatternFill(start_color="2B6CB0", end_color="2B6CB0", fill_type="solid")
            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            for col_idx, h in enumerate(self.headers, 1):
                cell = ws.cell(row=2, column=col_idx, value=h)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[2].height = 24
            
            # Widths (Col 6 is Defect Image)
            col_widths = [14, 20, 16, 16, 32, 22, 12, 12, 14, 12, 12, 12, 12, 18, 18]
            for idx, w in enumerate(col_widths, 1):
                ws.column_dimensions[openpyxl.utils.get_column_letter(idx)].width = w
                
            wb.save(self.filepath)
        elif HAS_OPENPYXL and self.filepath.exists():
            try:
                wb = openpyxl.load_workbook(self.filepath)
                ws = wb.active
                for row in ws.iter_rows(min_row=3, values_only=True):
                    if row and row[0]:
                        self.total_inspections += 1
                        if len(row) > 3 and str(row[3]).strip().upper() == "PASS":
                            self.total_pass += 1
                        else:
                            self.total_fail += 1
            except Exception:
                pass

        if not self.csv_path.exists():
            with open(self.csv_path, mode="w", newline="", encoding="utf-8") as f:
                writer = csv.writer(f)
                writer.writerow(self.headers)

    def log_inspection(self, device_name: str, res: dict, display_frame=None) -> dict:
        self.total_inspections += 1
        status = res.get("status", "FAIL")
        if status == "PASS":
            self.total_pass += 1
        else:
            self.total_fail += 1

        feats = res.get("features", {})
        min_conf = min([f.get("prob", 0.0) for f in feats.values()]) if feats else 0.0
        time_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        # Compile detailed failure reason
        defects = []
        for feat_name, f_info in feats.items():
            if not f_info.get("present", False):
                prob = f_info.get("prob", 0.0) * 100
                defects.append(f"{feat_name.replace('_', ' ').title()} ({prob:.0f}%)")
        
        fail_reason = "None - All features OK" if status == "PASS" else ("Missing: " + ", ".join(defects) if defects else "Alignment / Feature fault")

        # Handle defect image capture if FAIL
        defect_img_path = None
        if status != "PASS" and display_frame is not None:
            time_tag = datetime.now().strftime("%H%M%S")
            img_filename = f"defect_{self.total_inspections:04d}_{time_tag}.jpg"
            defect_img_path = self.defect_img_dir / img_filename
            # Resize image to compact thumbnail for fast Excel rendering
            thumb = cv2.resize(display_frame, (160, 160), interpolation=cv2.INTER_AREA)
            cv2.imwrite(str(defect_img_path), thumb)

        row = [
            f"#{self.total_inspections:04d}",
            time_str,
            device_name,
            status,
            fail_reason,
            "Photo Attached" if defect_img_path else "-",
            "PASS" if feats.get("spline", {}).get("present") else "FAIL",
            "PASS" if feats.get("top_nib", {}).get("present") else "FAIL",
            "PASS" if feats.get("bottom_nib", {}).get("present") else "FAIL",
            "PASS" if feats.get("hole_tl", {}).get("present") else "FAIL",
            "PASS" if feats.get("hole_tr", {}).get("present") else "FAIL",
            "PASS" if feats.get("hole_bl", {}).get("present") else "FAIL",
            "PASS" if feats.get("hole_br", {}).get("present") else "FAIL",
            round(min_conf * 100, 1),
            round(res.get("align_quality", {}).get("reproj_err", 0.0), 2)
        ]

        if HAS_OPENPYXL:
            try:
                wb = openpyxl.load_workbook(self.filepath)
                ws = wb.active
                target_row = ws.max_row + 1

                pass_fill = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
                fail_fill = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
                pass_font = Font(name="Calibri", size=11, bold=True, color="155724")
                fail_font = Font(name="Calibri", size=11, bold=True, color="721C24")

                # If defective image exists, adjust row height so thumbnail fits cleanly
                if defect_img_path:
                    ws.row_dimensions[target_row].height = 85
                else:
                    ws.row_dimensions[target_row].height = 22

                for col_idx, val in enumerate(row, 1):
                    cell = ws.cell(row=target_row, column=col_idx, value=val)
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    
                    # Overall Result (Column 4)
                    if col_idx == 4:
                        cell.fill = pass_fill if val == "PASS" else fail_fill
                        cell.font = pass_font if val == "PASS" else fail_font
                    # Failure Reason (Column 5)
                    elif col_idx == 5:
                        cell.alignment = Alignment(horizontal="left", vertical="center")
                        if status != "PASS":
                            cell.font = fail_font
                    # Defect Image (Column 6)
                    elif col_idx == 6 and defect_img_path:
                        cell.value = ""  # Cell text replaced by the embedded image
                    # Feature columns (Columns 7 to 13)
                    elif 7 <= col_idx <= 13:
                        if val == "FAIL":
                            cell.fill = fail_fill
                            cell.font = fail_font
                        else:
                            cell.fill = pass_fill
                            cell.font = pass_font

                # Insert the defect image into Column F (Col 6)
                if defect_img_path and defect_img_path.exists():
                    img = OpenpyxlImage(str(defect_img_path))
                    img.width = 100
                    img.height = 100
                    img_cell = f"F{target_row}"
                    ws.add_image(img, img_cell)

                wb.save(self.filepath)
            except Exception as e:
                print(f"[DailyExcelLogger] Excel write notice: {e}")

        try:
            with open(self.csv_path, mode="a", newline="", encoding="utf-8") as f:
                writer = csv.writer(f)
                writer.writerow(row)
        except Exception:
            pass

        return {
            "total_pass": self.total_pass,
            "total_fail": self.total_fail,
            "total_inspections": self.total_inspections
        }
